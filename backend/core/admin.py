from django.contrib import admin, messages
from django.contrib.auth.admin import UserAdmin
from django.http import HttpResponse
from django.urls import path
from django.shortcuts import render, redirect
from django import forms

from .models import (
    User,
    StudentProfile,
    QuestionTrigger,
    Question,
    Test,
    Attempt,
    Answer,
    SummaryReport,
    TestingPeriod,
)
from .trigger_visibility import get_public_trigger_for_source, public_triggers_queryset

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

import csv
import io

class CsvImportForm(forms.Form):
    csv_file = forms.FileField(label='CSV файл с русским переводом')


@admin.register(User)
class CustomUserAdmin(UserAdmin):
    list_display = [
        'username',
        'email',
        'role',
        'is_main_teacher',
        'training_mode',
        'created_by_teacher',
        'first_name',
        'last_name',
        'is_active',
        'is_staff',
    ]

    list_filter = [
        'role',
        'is_main_teacher',
        'training_mode',
        'created_by_teacher',
        'is_active',
        'is_staff',
    ]

    fieldsets = UserAdmin.fieldsets + (
        ('Role', {'fields': ('role', 'is_main_teacher', 'training_mode', 'created_by_teacher')}),
    )

    add_fieldsets = UserAdmin.add_fieldsets + (
        ('Role', {'fields': ('role', 'is_main_teacher', 'training_mode', 'created_by_teacher')}),
    )


@admin.register(StudentProfile)
class StudentProfileAdmin(admin.ModelAdmin):
    list_display = ['user', 'teacher', 'created_at']
    search_fields = ['user__username', 'user__email', 'user__created_by_teacher__username']

    def teacher(self, obj):
        return obj.user.created_by_teacher


@admin.register(QuestionTrigger)
class QuestionTriggerAdmin(admin.ModelAdmin):
    list_display = ['name', 'slug', 'is_active', 'question_count', 'created_at']
    list_filter = ['is_active']
    search_fields = ['name', 'slug', 'description']
    prepopulated_fields = {'slug': ('name',)}

    def question_count(self, obj):
        return obj.questions.count()

    question_count.short_description = 'Question count'


@admin.register(Question)
class QuestionAdmin(admin.ModelAdmin):
    list_display = [
        'id',
        'text_preview',
        'text_ru',
        'public_trigger',
        'correct_answer',
        'created_by',
        'created_at',
    ]

    list_filter = ['correct_answer', 'created_by']

    search_fields = [
        'id',
        'text',
        'text_ru',
        'text_ky',
        'option_a',
        'option_a_ru',
        'option_a_ky',
        'option_b',
        'option_b_ru',
        'option_b_ky',
        'option_c',
        'option_c_ru',
        'option_c_ky',
        'option_d',
        'option_d_ru',
        'option_d_ky',
    ]

    change_list_template = 'admin/questions_changelist.html'

    fieldsets = (
        ('Основная информация', {
            'fields': (
                'text',
                'text_ru',
                'text_ky',
                'trigger',
                'correct_answer',
                'created_by',
            )
        }),
        ('Варианты ответов', {
            'fields': (
                'option_a',
                'option_a_ru',
                'option_a_ky',
                'option_b',
                'option_b_ru',
                'option_b_ky',
                'option_c',
                'option_c_ru',
                'option_c_ky',
                'option_d',
                'option_d_ru',
                'option_d_ky',
            )
        }),
        ('Дополнительно', {
            'fields': (
                'explanation',
            )
        }),
    )

    def get_urls(self):
        urls = super().get_urls()

        custom_urls = [
            path(
                'import-ru-translations/',
                self.admin_site.admin_view(self.import_ru_translations),
                name='import_ru_translations',
            ),
        ]

        return custom_urls + urls

    def import_ru_translations(self, request):
        if request.method == 'POST':
            form = CsvImportForm(request.POST, request.FILES)

            if form.is_valid():
                csv_file = request.FILES['csv_file']

                try:
                    decoded_file = csv_file.read().decode('utf-8-sig')
                except UnicodeDecodeError:
                    messages.error(request, 'Ошибка кодировки файла. Сохрани CSV в UTF-8.')
                    return redirect('.')

                reader = csv.DictReader(io.StringIO(decoded_file))

                updated = 0
                not_found = 0
                skipped = 0

                for row in reader:
                    question_id = row.get('id')

                    if not question_id:
                        skipped += 1
                        continue

                    try:
                        question = Question.objects.get(id=int(question_id))
                    except Question.DoesNotExist:
                        not_found += 1
                        continue
                    except ValueError:
                        skipped += 1
                        continue

                    question.text_ru = row.get('text_ru', '').strip()
                    question.option_a_ru = row.get('option_a_ru', '').strip()
                    question.option_b_ru = row.get('option_b_ru', '').strip()
                    question.option_c_ru = row.get('option_c_ru', '').strip()
                    question.option_d_ru = row.get('option_d_ru', '').strip()

                    question.save()
                    updated += 1

                messages.success(
                    request,
                    f'Импорт завершён. Обновлено: {updated}. Не найдено: {not_found}. Пропущено: {skipped}.'
                )

                return redirect('..')

        else:
            form = CsvImportForm()

        return render(request, 'admin/import_ru_translations.html', {'form': form})

    def text_preview(self, obj):
        return obj.text[:50] + '...' if len(obj.text) > 50 else obj.text

    def public_trigger(self, obj):
        trigger = get_public_trigger_for_source(obj.trigger, QuestionTrigger)
        return trigger.name if trigger else '-'

    public_trigger.short_description = 'Триггер'


@admin.register(Test)
class TestAdmin(admin.ModelAdmin):
    list_display = [
        'title',
        'teacher',
        'language',
        'is_translation_from_ky',
        'question_count',
        'is_active',
        'created_at',
    ]

    list_filter = ['is_active', 'teacher', 'language', 'is_translation_from_ky']
    search_fields = ['title']
    filter_horizontal = ['questions', 'triggers']
    actions = ['export_to_word_russian']

    def question_count(self, obj):
        return obj.questions.count()

    def formfield_for_manytomany(self, db_field, request, **kwargs):
        if db_field.name == 'triggers':
            kwargs['queryset'] = public_triggers_queryset(QuestionTrigger)

        return super().formfield_for_manytomany(db_field, request, **kwargs)

    def export_to_word_russian(self, request, queryset):
        for test in queryset:
            doc = Document()

            title = doc.add_heading(test.title, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph(f'Язык: {test.get_language_display()}')

            if test.is_translation_from_ky:
                doc.add_paragraph('Перевод с кыргызского на русский: Да')
            else:
                doc.add_paragraph('Перевод с кыргызского на русский: Нет')

            doc.add_paragraph(f'Экзаменатор: {test.teacher.username}')
            doc.add_paragraph(f'Количество вопросов: {test.questions.count()}')
            doc.add_paragraph('')

            for idx, question in enumerate(test.questions.all().order_by('id'), 1):
                q_text = question.text_ru if question.text_ru else question.text

                doc.add_heading(f'Вопрос {idx}', level=1)
                doc.add_paragraph(q_text)

                options = [
                    ('A', question.option_a_ru if question.option_a_ru else question.option_a),
                    ('B', question.option_b_ru if question.option_b_ru else question.option_b),
                    ('C', question.option_c_ru if question.option_c_ru else question.option_c),
                    ('D', question.option_d_ru if question.option_d_ru else question.option_d),
                ]

                for opt_label, opt_text in options:
                    if opt_text:
                        p = doc.add_paragraph(f'{opt_label}) {opt_text}')
                        p.paragraph_format.left_indent = Inches(0.25)

                doc.add_paragraph(f'Правильный ответ: {question.correct_answer}')
                doc.add_paragraph('')

            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )

            filename = f'{test.title}_russian.docx'
            response['Content-Disposition'] = f'attachment; filename="{filename}"'

            return response

    export_to_word_russian.short_description = 'Экспорт в Word (русский)'


@admin.register(Attempt)
class AttemptAdmin(admin.ModelAdmin):
    list_display = [
        'student',
        'test',
        'status',
        'score',
        'started_at',
        'completed_at',
        'archived_at',
    ]

    list_filter = ['status', 'test', 'started_at', 'archived_at']
    search_fields = ['student__username', 'test__title']


@admin.register(Answer)
class AnswerAdmin(admin.ModelAdmin):
    list_display = [
        'attempt',
        'question',
        'selected_option',
        'is_correct',
        'answered_at',
    ]

    list_filter = ['is_correct', 'selected_option']


@admin.register(SummaryReport)
class SummaryReportAdmin(admin.ModelAdmin):
    list_display = [
        'title',
        'teacher',
        'testing_period',
        'student_count',
        'generated_at',
        'expires_at',
    ]

    list_filter = ['teacher', 'generated_at', 'expires_at']
    search_fields = [
        'title',
        'teacher__username',
        'teacher__first_name',
        'teacher__last_name',
    ]

    readonly_fields = ['generated_at']


@admin.register(TestingPeriod)
class TestingPeriodAdmin(admin.ModelAdmin):
    list_display = [
        'id',
        'teacher',
        'started_at',
        'ends_at',
        'summary_delete_at',
        'processed_at',
        'summary_created_at',
    ]

    list_filter = ['teacher', 'processed_at']
    search_fields = [
        'teacher__username',
        'teacher__first_name',
        'teacher__last_name',
    ]

    readonly_fields = ['created_at', 'processed_at', 'summary_created_at']
